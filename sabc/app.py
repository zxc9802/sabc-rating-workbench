from io import BytesIO
from copy import deepcopy
import json
import os
from pathlib import Path
from uuid import uuid4
from threading import Lock
import zipfile

from starlette.concurrency import run_in_threadpool
from fastapi import FastAPI, HTTPException, UploadFile, Form
from fastapi.responses import JSONResponse, Response
from pydantic import BaseModel, Field

from sabc.catalog import catalog
from sabc.key_storage import encrypt_key, model_key, key_origin
from sabc.request_limits import BodyLimitMiddleware
from sabc.llm import analyze, guide
from sabc import planner, report_qa, report_readiness, advisory, framing, report_grounding
from sabc.fact_boundaries import gap_details
from sabc import auth, lifecycle, model_router, sso
from sabc.tenancy import AccountStore, account_id
from sabc import speech
from sabc.jobs import jobs
from sabc.streaming import check_cancelled, progress
from sabc.rating import assess, DIMENSIONS, TYPES, RULE_VERSION, PROJECT_FIELDS
from sabc.store import Store, utcnow
from sabc.schema import validate_amounts, validate_proposal, validate_project_type
from sabc.sources import collect, SUPPORTED
from sabc.local_sources import REGIONS, save_capture, use_capture

ROOT=Path(__file__).resolve().parent.parent
store=AccountStore(Path(os.getenv('SABC_DB',str(ROOT/'data'/'sabc.db'))))
app=FastAPI(title='SABC 项目评级',docs_url=None,redoc_url=None,openapi_url=None)
app.add_middleware(BodyLimitMiddleware)
app.include_router(speech.router)
app.include_router(auth.router)
app.include_router(sso.router)
initial_interview_lock=Lock()


@app.middleware('http')
async def local_writes(request,call_next):
    user = None
    public = request.url.path in ('/api/auth/session','/api/auth/login','/api/health','/api/sso/start','/api/sso/callback')
    if sso.enabled() and not public:
        try:
            user = await run_in_threadpool(sso.identity, request)
        except HTTPException as error:
            return JSONResponse({'detail':error.detail},status_code=error.status_code)
        if not user:
            return JSONResponse({'detail':'请通过主站登录后继续'},status_code=401)
    elif not public and not auth.authenticated(request):
        return JSONResponse({'detail':'请登录后继续'},status_code=401,headers={'Cache-Control':'no-store'})
    origin=request.headers.get('origin')
    if request.method not in ('GET','HEAD','OPTIONS') and origin and origin not in ('http://127.0.0.1:3000','http://localhost:3000','http://127.0.0.1:18765',os.getenv('SABC_UI_ORIGIN','http://127.0.0.1:3000')):
        return JSONResponse({'detail':'不接受其他网站的写入请求'},status_code=403)
    token = account_id.set(user['id'] if user else None)
    try:
        response=await call_next(request)
    finally:
        account_id.reset(token)
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['Cache-Control']='no-store'
    return response


@app.exception_handler(ValueError)
async def invalid(request,error):
    return JSONResponse({'detail':str(error)},status_code=422)


def project_or_404(pid):
    p=store.get('projects',pid)
    if not p: raise HTTPException(404,'项目不存在')
    return p


def company():
    records=store.list('companies')
    return records[0] if records else {}


def evidence_for(pid):
    return [e for e in store.list('evidence') if e.get('project_id')==pid]


def model_evidence_for(pid):
    return [e for e in evidence_for(pid) if not e.get('capture_id') and e.get('capture_method')!='browser-observation']


def settings():
    return store.get('settings','model') or {'base_url':os.getenv('SABC_MODEL_BASE_URL',''),'model':os.getenv('SABC_MODEL',''),'encrypted_key':''}


@app.get('/api/health')
def health():
    return {'status':'ok'}


def public_settings():
    s=settings()
    primary=model_router.deepseek()
    mixtoken=model_router.mixtoken_route()
    gemini=model_router.gemini_route({**s, 'key': bool(not s.get('key_disabled') and (s.get('encrypted_key') or os.getenv('SABC_API_KEY')))})
    result = {'managed':sso.enabled(),'primary_model':mixtoken['model'] if mixtoken else gemini['model'] if gemini else ('z-ai/glm-5.3-flash' if os.getenv('SABC_FAL_API_KEY','').strip() else s.get('model','')), 'planner_model':'八维程序规则', 'reasoning_effort':primary['effort'] if primary else None, 'fallback_model':primary['model'] if primary else '', 'base_url':s.get('base_url',''),'model':s.get('model',''),
            'has_key':bool(mixtoken or primary or (not s.get('key_disabled') and (s.get('encrypted_key') or os.getenv('SABC_API_KEY')))),
            'configured':bool(mixtoken or primary or (s.get('base_url') and s.get('model')))}
    if sso.enabled():
        result.update(base_url='', model='', primary_model='项目分析模型', fallback_model='', reasoning_effort=None)
    return result


@app.get('/api/bootstrap')
def bootstrap():
    sources=catalog()
    runs=store.list('source_runs')
    for source in sources:
        source['status']='connected' if any(r['source']==source['id'] and r['status']=='success' for r in runs) else 'available' if source['id'] in SUPPORTED else 'browser' if source['id'] in ('baidu','douyin') else 'pending'
    return {'projects':[{**p,'followup':lifecycle.followup(p)} for p in store.list('projects')],'company':company(),'settings':public_settings(),
            'sources':sources,'rule_version':RULE_VERSION,'types':TYPES,
            'dimensions':[{'key':k,'name':n,'weight':w} for k,(n,w) in DIMENSIONS.items()]}


@app.put('/api/company')
def save_company(body:dict):
    allowed={'name','strategy','team','budget','cash_available','cash_safety_line','confirmed','capabilities','active_projects','risk_policy','approved_by'}
    data={k:v for k,v in body.items() if k in allowed}
    if 'confirmed' in data and not isinstance(data['confirmed'],bool):
        raise ValueError('公司确认状态必须为布尔值')
    validate_amounts(data, ('budget','cash_available','cash_safety_line'))
    for key in ('budget','cash_available','cash_safety_line'):
        if data.get(key) is not None and (not isinstance(data[key],(int,float)) or isinstance(data[key],bool) or data[key]<0):
            raise ValueError('预算和现金必须为非负数')
    data['version']=len(store.list('companies'))+1
    return store.save('companies',data)


@app.post('/api/projects')
def create_project(body:dict):
    validate_project_type(body)
    validate_amounts(body, ('budget_requested',))
    data={k:v for k,v in body.items() if k in set(PROJECT_FIELDS)|{'description','budget_requested'}}
    data['name']=str(data.get('name') or '未命名项目')[:100]
    data['project_type']=data.get('project_type')
    data['version']=1
    data['messages']=[]
    data['lifecycle']={**lifecycle.initial(), 'mode': 'continuous'}
    if body.get('stage') and body['stage']!='pre':
        data=lifecycle.transition(data,'set_stage',body,company())
    saved=store.save('projects',data)
    if body.get('auto_start'):
        try:
            start_interview(saved['id'])
        except HTTPException as error:
            # Creation succeeded; a busy interview queue must not invite duplicate projects.
            if error.status_code != 429: raise
        saved=project_or_404(saved['id'])
    return saved


class DeleteProjects(BaseModel):
    ids:list[str]=Field(min_length=1,max_length=200)


@app.post('/api/projects/delete')
def delete_projects(body:DeleteProjects):
    ids=list(dict.fromkeys(body.ids))
    with jobs.lock:
        for job in store.list('jobs'):
            if job.get('project_id') in ids and job.get('status')=='running':
                jobs.cancel_locked(store,job['id'])
        store.delete_projects(ids)
    return {'deleted':ids}


@app.get('/api/projects/{pid}')
def get_project(pid:str):
    completed = [j['id'] for j in store.list('jobs') if j.get('project_id') == pid and j['status'] != 'running']
    p=project_or_404(pid)
    evidence=evidence_for(pid)
    analysis_complete=report_readiness.ready(p,company(),evidence)
    report_ready=analysis_complete
    return {'project':{**p,'followup':lifecycle.followup(p),'analysis_complete':analysis_complete, 'report_ready':report_ready, 'report_pipeline':pipeline_status(p,company(),evidence)},'evidence':evidence,
            'assessments':[r for r in store.list('assessments') if r['project_id']==pid], 'completed_job_ids':completed,
            'active_jobs':[jobs.read(store,r['id']) for r in store.list('jobs') if r['project_id']==pid and r['status']=='running' and r.get('operation')!='report_chat'],
            'initial_job':jobs.read(store,p['initial_job_id']) if p.get('initial_job_id') else None}


@app.post('/api/projects/{pid}/start-interview')
def start_interview(pid:str,retry:bool=False):
    with initial_interview_lock:
        p=project_or_404(pid)
        if p['messages']: return {'status':'success','id':''}
        existing=next((j for j in store.list('jobs') if j['project_id']==pid and (j['id']==p.get('initial_job_id') or j.get('operation')=='chat')),None)
        if existing:
            existing=jobs.read(store,existing['id'])
            if not retry or existing['status']=='running': return existing
        ident=str(uuid4())
        message=str(p.get('description') or '请结合公司资料开始项目访谈。')[:12000]
        return jobs.submit(store,ident,pid,{'operation':'chat','payload':{'message':message}},lambda:chat(pid,Chat(message=message)),queue=True)


@app.patch('/api/projects/{pid}')
def update_project(pid:str,body:dict):
    p=project_or_404(pid)
    validate_project_type(body)
    allowed=set(PROJECT_FIELDS)|{'description','budget_requested','proposal'}
    patch={k:v for k,v in body.items() if k in allowed}
    validate_amounts(patch, ('budget_requested',))
    if patch.get('proposal') is not None: patch['proposal']=validate_proposal(patch['proposal'])
    return store.save('projects',{**p,**patch,'pending_patch':{},'version':p['version']+1,
                                 'risks_source':'user' if 'risks' in patch else p.get('risks_source', 'unknown')})



class LifecycleAction(BaseModel):
    action:str
    payload:dict=Field(default_factory=dict)
    version:int


@app.post('/api/projects/{pid}/lifecycle')
def update_lifecycle(pid:str,body:LifecycleAction):
    with jobs.lock:
        p=project_or_404(pid)
        if p['version']!=body.version:
            raise HTTPException(409,'项目已更新，请刷新后再操作')
        for job in store.list('jobs'):
            if job.get('project_id')==pid and jobs.read(store,job['id'])['status']=='running':
                raise HTTPException(409,'请等待本项目当前任务完成后再修改阶段')
        previous = None
        if body.action == 'advance':
            stage = lifecycle.state(p)['stage']
            previous = next((a for a in store.list('assessments') if a.get('project_id') == pid
                             and (a['result'].get('stage') or a['snapshot']['project'].get('lifecycle', {}).get('stage')) == stage), None)
            if not previous:
                raise ValueError('请先生成并保存当前阶段报告，再进入下一阶段')
        changed=lifecycle.transition(p,body.action,body.payload,company())
        if previous:
            changed['lifecycle']['previous_report_id'] = previous['id']
        if body.action in ('set_stage','confirm_plan','start','complete','advance'):
            changed['proposal']=None
            changed['interview']={}
            # Old ratings remain in immutable reports, never label a new stage as final.
            changed.pop('last_grade',None)
        return store.save('projects',changed)


@app.get('/api/projects/{pid}/model-runs')
def model_runs(pid:str):
    project_or_404(pid)
    runs = [r for r in store.list('model_runs') if r.get('project_id')==pid][:100]
    return [{k:v for k,v in r.items() if k not in ('provider','model')} for r in runs] if sso.enabled() else runs


class Chat(BaseModel):
    revision_of: str | None = None
    revision_turn_id: str | None = None
    message:str=Field(min_length=1,max_length=12000)
    field:str|None=None
    generate_report:bool=False


@app.post('/api/projects/{pid}/chat')
def chat_request(pid:str,body:Chat):
    project_or_404(pid)
    return jobs.execute(store,pid,lambda:chat(pid,body))


def chat(pid:str,body:Chat):
    token=model_router.audit.set(lambda event:store.save('model_runs',{'project_id':pid,**event}))
    notify = progress.get()
    # Interview JSON is normalized after SSE completes; never expose a draft
    # question that may be rejected or replaced by the collection gate.
    display_progress = progress.set((lambda _: notify('正在生成报告…')) if body.generate_report else (lambda _: None)) if notify else None
    try:
        result = chat_turn(pid,body)
        if notify and not body.generate_report:
            notify(result['reply'])
        return result
    finally:
        if display_progress is not None: progress.reset(display_progress)
        model_router.audit.reset(token)


def chat_turn(pid,body):
    check_cancelled()
    p=project_or_404(pid)
    previous_report = None
    previous = next((a for a in store.list('assessments') if a.get('project_id') == pid), None)
    if previous:
        previous_report = {'id': previous['id'], 'created_at': previous['created_at'], 'result': previous['result'],
                           'lifecycle': lifecycle.context(previous['snapshot']['project'])}
    c=company(); e=evidence_for(pid)
    if body.revision_of:
        report_or_404(pid, body.revision_of)
        turn = store.get('report_turns', body.revision_turn_id or '')
        if not turn or turn.get('project_id') != pid or turn.get('assessment_id') != body.revision_of:
            raise ValueError('请选择当前项目已保存的报告问答用于修订')
        if not lifecycle.collection_ready(p.get('lifecycle', {})) or p.get('interview', {}).get('questions'):
            return {'mode':'model','needs_collection':True,'reply':'当前仍有待回答问题，请先完成本轮问答再修订报告。'}
        revision={'assessment_id':body.revision_of,'turn_id':body.revision_turn_id,'question':turn['question']}
        if p.get('report_revision') != revision:
            p['report_revision']=revision
            p['version']+=1
            p['collection_completion']={'input_fingerprint':report_readiness.fingerprint(p,c,e),'collection_version':4}
            store.save('projects',p)
        return generate_final_report(p,c,e,previous_report)
    workflow=current_workflow(p,c,e)
    same_answer = next((m.get('content') for m in reversed(p.get('messages', [])) if m.get('role')=='user'), None)==body.message
    resumable = workflow and workflow.get('input_fingerprint')==report_readiness.fingerprint(p,c,e)
    if body.generate_report or (resumable and same_answer):
        if not report_readiness.ready(p,c,e):
            return {'mode':'model','needs_collection':True,'reply':'资料发生变化，请先继续问答，更新本次判断。'}
        return generate_final_report(p,c,e,previous_report)
    p['lifecycle'] = {**lifecycle.state(p), 'mode': 'continuous', 'confirmed': True}
    p.pop('assessment_review', None)
    p.pop('report_preparation', None)
    p.pop('collection_completion', None)
    p.pop('report_revision', None)
    p['proposal']=None
    # New input invalidates the old approval even if this interview request fails.
    with jobs.lock:
        check_cancelled()
        store.save('projects',p)
    messages=list(p.get('messages',[]))
    # A failed internal call has already saved the user's answer. Retry that turn.
    if not messages or messages[-1].get('role') != 'user' or messages[-1].get('content') != body.message:
        messages.append({'role':'user','content':body.message,'stage':lifecycle.state(p)['stage'],'time':utcnow()})
    s=settings()
    if model_router.mixtoken_route() or model_router.deepseek() or (s.get('base_url') and s.get('model')):
        from sabc.dimension_sources import rule_plan
        from concurrent.futures import ThreadPoolExecutor
        from contextvars import copy_context
        plan=rule_plan(p,body.message,model_evidence_for(pid))
        store.save('retrieval_plans',{'project_id':pid,'created_at':utcnow(),**plan})
        def retrieve(request):
            check_cancelled()
            try:
                evidence=collect(store,pid,request['source'],request['query'])
                return {**request,'evidence_id':evidence['id'],'status':'saved'}
            except Exception as error:
                store.save('source_runs',{'project_id':pid,'source':request['source'],'query':request['query'],'status':'failed','error':str(error),'boundary':'chat'})
                return {**request,'status':'failed','reason':'外部核查未完成'}
        outcomes=[]
        if plan['data_requests']:
            with ThreadPoolExecutor(max_workers=2) as pool:
                futures=[pool.submit(copy_context().run,retrieve,r) for r in plan['data_requests']]
                outcomes=[f.result() for f in futures]
        check_cancelled()
        context={'results':outcomes,'candidates':plan['candidates'],'missing_parameters':plan['missing_parameters']}
        starting_inputs=report_readiness.fingerprint(p,company(),evidence_for(pid))
        result=analyze(s,model_key(s),{**p, '_report_requested': False, '_previous_stage_report': previous_report},company(),model_evidence_for(pid),messages+[{'role':'tool','content':'程序已完成本轮规则取数。saved仅表示已保存而非已核验；failed表示未完成，不得虚构结果。缺少参数时先追问。仅依据现有证据回答，不再请求取数：'+json.dumps(context,ensure_ascii=False)}])
        result['proposal']=None
        result['stage_review']=None
        result['pilot_plan']=None
        result['retrieval_results']=outcomes
        result['retrieval_plan']=plan
        result['data_requests']=[]
        # Keep earlier unconfirmed facts until the user accepts them; later explicit
        # corrections replace the same field, not the whole pending fact set.
        p['pending_patch']={**p.get('pending_patch',{}),**result.get('project_patch',{})}
        framing.apply(p, result.get('framing'), messages)
    else:
        result=guide(p,body.message,body.field)
        p.update(result['project_patch'])
    report_grounding.preserve_period(p)
    followups=[q for q in result.get('questions',[])[:2] if q.strip() and q not in result['reply']]
    # Some providers put the questions in both fields, with different wording.
    # Keep the intact conversational reply instead of adding a second interview.
    if followups and not any(mark in result['reply'] for mark in ('？','?')):
        result['reply']+='\n\n'+'\n\n'.join(followups)
    valid_refs={e['id'] for e in model_evidence_for(pid)}
    refs=[eid for eid in result.get('reply_evidence_ids',[]) if eid in valid_refs]
    p['messages']=messages+[{'role':'assistant','content':result['reply'],'mode':result['mode'],'field':result.get('field'),'stage':lifecycle.state(p)['stage'],'evidence_ids':refs,'time':utcnow(),
                            'question_targets':result.get('question_targets', [])}]
    if result['mode']=='model': p['proposal']=result.get('proposal')
    if result['mode']=='model':
        lifecycle.absorb(p,result,company(),model_evidence_for(pid))
        life = p.get('lifecycle', {})
        gaps = lifecycle.collection_gaps(life)
        state='gathering' if result.get('questions') else 'ready' if lifecycle.collection_ready(life) else 'paused'
        p['interview']={'state':state,'gaps':gaps,'questions':[] if state!='gathering' else result.get('questions',[]),
                        'question_targets':[] if state!='gathering' else result.get('question_targets',[]),
                        'note':'第一阶段问答已完成' if state=='ready' else '可继续补充资料或讨论下一步验证办法' if state=='paused' else '补充影响决策的关键事实'}
    if result['mode']=='model' and not result.get('questions') and lifecycle.collection_ready(p.get('lifecycle', {})):
        # Commit the interview before drafting. The draft lives in a private record.
        p['messages'] = messages + [{'role':'assistant','content':'第一阶段问答已完成，正在整理报告。',
                                    'mode':'model','stage':lifecycle.state(p)['stage'],'evidence_ids':[],'time':utcnow()}]
        p['interview'].update(state='ready', questions=[], note='第一阶段问答已完成')
        c=company(); e=evidence_for(pid)
        p['collection_completion']={'input_fingerprint':report_readiness.fingerprint(p,c,e), 'collection_version':4}
        p['version']+=1
        with jobs.lock:
            check_cancelled()
            if report_readiness.fingerprint(project_or_404(pid), c, e) != starting_inputs:
                raise ValueError('分析期间项目资料发生变化，请根据最新资料继续。')
            store.save('projects', p)
        return generate_final_report(p,c,e,previous_report)
    p['version']+=1
    with jobs.lock:
        check_cancelled()
        if result['mode']=='model':
            current=project_or_404(pid)
            if report_readiness.fingerprint(current,company(),evidence_for(pid))!=starting_inputs:
                raise ValueError('分析期间项目资料发生变化，请根据最新资料继续。')
        store.save('projects',p)
    return result


def current_workflow(project, c, e):
    workflow=store.get('report_workflows',project['id'])
    if not workflow:
        return None
    current_inputs=report_readiness.fingerprint(project,c,e)
    if workflow.get('input_fingerprint')==current_inputs:
        return workflow
    # Recover a crash after final project persistence but before workflow completion.
    if workflow.get('final_id') and workflow['final_id']==project.get('last_assessment_id'):
        record=store.get('assessments',workflow['final_id'])
        if record and report_readiness.fingerprint(record['snapshot']['project'],record['snapshot']['company'],record['snapshot']['evidence'])==current_inputs:
            return {**workflow,'step':'complete','report_id':record['id'],'input_fingerprint':current_inputs}
    return None


def pipeline_status(project, c, e):
    workflow=current_workflow(project,c,e)
    # This is the entire public contract; candidate and review notes stay private.
    return {'step':workflow['step'], 'report_id':workflow.get('report_id')} if workflow else None


def generate_final_report(p,c,e,previous_report):
    from sabc.report_corrections import ReportCorrections
    pid=p['id']
    starting_inputs=report_readiness.fingerprint(p,c,e)
    workflow=current_workflow(p,c,e)
    if not workflow:
        workflow={'id':pid,'project_id':pid,'input_fingerprint':starting_inputs,'step':'preparing'}
    if workflow.get('report_id'):
        return {'mode':'model','report_id':workflow['report_id'],'reply':'最终报告已生成。'}
    correction_state=workflow.setdefault('corrections',{'count':len(workflow.get('notes',[]))})

    def checkpoint(candidate=None, notes=None):
        nonlocal workflow
        with jobs.lock:
            check_cancelled()
            if report_readiness.fingerprint(project_or_404(pid),company(),evidence_for(pid))!=starting_inputs:
                raise ValueError('处理期间资料发生变化，请根据最新资料继续问答。')
            if candidate is not None:
                workflow.update(candidate=candidate, notes=notes or [], step='revising' if notes else 'reviewing')
            workflow['corrections']=correction_state
            workflow=store.save('report_workflows',workflow)

    checkpoint()
    s={**settings(),'report_corrections':ReportCorrections(correction_state,checkpoint)}
    if not workflow.get('candidate'):
        result=analyze(s,model_key(s),{**p,'_report_requested':True,'_previous_stage_report':previous_report},
                       c,model_evidence_for(pid),p.get('messages',[]))
        check_cancelled()
        if not result.get('proposal') or not result.get('stage_review') or result.get('questions'):
            raise ValueError('报告整理未完成，请继续处理；已收集的资料仍然保留。')
        accepted=deepcopy(p)
        patch=accepted.pop('pending_patch',{})
        accepted.update(patch)
        accepted['pending_patch']={}
        if 'risks' in patch: accepted['risks_source']='model'
        accepted['version']+=bool(patch)
        accepted['proposal']=validate_proposal(result['proposal'])
        lifecycle.absorb(accepted,{'proposal':accepted['proposal'],'stage_review':result['stage_review']},c,e)
        checkpoint(build_assessment(accepted,c,e,accepted['proposal']))
    candidate=advisory.review_report(s,model_key(s),workflow['candidate'],build_assessment,
                                     notes=workflow.get('notes'),checkpoint=checkpoint)
    # Save a stable report ID before publishing. A retry cannot create a duplicate
    # even if the process stops between saving the report and updating the project.
    checkpoint(candidate, candidate['snapshot']['quality_review']['rounds'])
    with jobs.lock:
        check_cancelled()
        if report_readiness.fingerprint(project_or_404(pid),company(),evidence_for(pid))!=starting_inputs:
            raise ValueError('处理期间资料发生变化，请根据最新资料继续问答。')
        final_project=candidate['snapshot']['project']
        if p.get('report_revision'):
            old=report_or_404(pid,p['report_revision']['assessment_id'])
            changes=[]
            for key,(name,_) in DIMENSIONS.items():
                if old['snapshot']['proposal'].get('dimensions',{}).get(key) != candidate['snapshot']['proposal']['dimensions'].get(key):
                    changes.append(name+'的分数或依据')
            if old['snapshot']['proposal'].get('decision_brief') != candidate['snapshot']['proposal'].get('decision_brief'):
                changes.append('关键判断与升级说明')
            for field,label in (('assessment_scope','本次评估范围'),('pros','项目优势'),
                                ('cons','项目劣势'),('strongest_objections','最大隐患')):
                if old['snapshot']['proposal'].get(field) != candidate['snapshot']['proposal'].get(field):
                    changes.append(label)
            for field,label in (('timeframe','未来验证周期'),('data_period','历史资料期间')):
                if old['snapshot']['project'].get(field) != final_project.get(field):
                    changes.append(label)
            old_review=old['snapshot']['project'].get('lifecycle',{}).get('review') or {}
            new_review=final_project.get('lifecycle',{}).get('review') or {}
            if any(old_review.get(field)!=new_review.get(field) for field in ('summary','next_action','next_review_days','conclusion')):
                changes.append('行动建议')
            candidate['result']['revision']={'source_report_id':old['id'],'changes':changes}
        completed='最终报告已生成。' if candidate['snapshot']['quality_review'].get('status')=='revision_limit' else '第二阶段审查已完成，最终报告已生成。'
        final_project['messages'].append({'role':'assistant','content':completed,
                                         'mode':'model','stage':lifecycle.state(final_project)['stage'],'evidence_ids':[],'time':utcnow()})
        final_project['collection_completion']={'input_fingerprint':report_readiness.fingerprint(final_project,c,e),'collection_version':4}
        candidate['id']=workflow.setdefault('final_id',uuid4().hex)
        store.save('report_workflows',workflow)
        record=store.get('assessments',candidate['id'])
        if record:
            saved=record['snapshot']['project']
            store.save('projects',{**saved,'last_grade':record['result']['grade'],'last_assessment_id':record['id']})
        else:
            record=persist_assessment(candidate)
        store.save('report_workflows',{**workflow,'step':'complete','report_id':record['id'],
                   'input_fingerprint':report_readiness.fingerprint(record['snapshot']['project'],c,e)})
    return {'mode':'model','report_id':record['id'],'reply':'最终报告已生成。'}


class Evidence(BaseModel):
    publisher: str = Field(default='',max_length=200)
    title:str=Field(min_length=1,max_length=200)
    source_locator:str=Field(min_length=1,max_length=2000)
    content:str=Field(default='',max_length=200000)
    source_type:str='user'
    data_period:str=''
    scope:str=''
    verification_status:str='unverified'
    level:int=Field(default=0,ge=0,le=3)
    repeat_verified:bool=False
    valid_until:str|None=None
    conflict:bool=False
    canonical_source:str|None=None


class SourceQuery(BaseModel):
    query: str = Field(min_length=1,max_length=200)


@app.post('/api/projects/{pid}/sources/{source}')
def source_request(pid:str,source:str,body:SourceQuery):
    project_or_404(pid)
    return jobs.execute(store,pid,lambda:fetch_source(pid,source,body))


def fetch_source(pid:str,source:str,body:SourceQuery):
    project_or_404(pid)
    try:
        return collect(store,pid,source,body.query)
    except Exception as error:
        store.save('source_runs',{'project_id':pid,'source':source,'query':body.query,'status':'failed','error':str(error),'boundary':'manual'})
        return {'status':'skipped'}


class ReportQuestion(BaseModel):
    assessment_id:str=Field(min_length=1,max_length=100)
    message:str=Field(min_length=1,max_length=6000)


def report_or_404(pid, ident):
    project_or_404(pid)
    report=store.get('assessments',ident)
    if not report or report['project_id']!=pid: raise HTTPException(404,'报告不存在')
    return report


def report_turns(pid, ident):
    return sorted([t for t in store.list('report_turns') if t['project_id']==pid and t['assessment_id']==ident],key=lambda t:t['created_at'])


@app.get('/api/projects/{pid}/reports/{ident}/chat')
def report_history(pid:str,ident:str):
    report_or_404(pid,ident)
    active=[jobs.read(store,j['id']) for j in store.list('jobs') if j['project_id']==pid and j.get('assessment_id')==ident and j['status']=='running']
    return {'turns':report_turns(pid,ident),'active_job':next((j for j in active if j['status']=='running'),None)}


def report_chat(pid, body, job_id):
    report=report_or_404(pid,body.assessment_id)
    config=settings()
    token=model_router.audit.set(lambda event:store.save('model_runs',{'project_id':pid,'assessment_id':body.assessment_id,**event}))
    try:
        reply=report_qa.answer(config,model_key(config),report,report_turns(pid,body.assessment_id),body.message.strip())
    finally:
        model_router.audit.reset(token)
    with jobs.lock:
        check_cancelled()
        return store.save('report_turns',{'id':job_id,'project_id':pid,'assessment_id':body.assessment_id,'question':body.message.strip(),'reply':reply})


class SlowOperation(BaseModel):
    id:str=Field(pattern=r'^[a-f0-9-]{36}$')
    operation:str=Field(pattern=r'^(chat|source|report_chat)$')
    source:str=''
    payload:dict


@app.post('/api/projects/{pid}/jobs',status_code=202)
def start_job(pid:str,body:SlowOperation):
    project_or_404(pid)
    if body.operation=='chat':
        parsed=Chat.model_validate(body.payload)
        action=lambda: chat(pid,parsed)
    elif body.operation=='report_chat':
        parsed=ReportQuestion.model_validate(body.payload)
        if not parsed.message.strip(): raise ValueError('请输入报告问题')
        report_or_404(pid,parsed.assessment_id)
        action=lambda: report_chat(pid,parsed,body.id)
    else:
        parsed=SourceQuery.model_validate(body.payload)
        from sabc.sources import request_spec
        request_spec(body.source,parsed.query)
        action=lambda: fetch_source(pid,body.source,parsed)
    return jobs.submit(store,body.id,pid,body.model_dump(exclude={'id'}),action)


@app.get('/api/jobs/{ident}')
def get_job(ident:str):
    job=jobs.read(store,ident)
    if job.get('operation')=='chat':
        p=project_or_404(job['project_id'])
        job={**job,'collection':{'lifecycle':p.get('lifecycle'), 'report_ready':report_readiness.ready(p,company(),evidence_for(p['id'])),
                                  'report_pipeline':pipeline_status(p,company(),evidence_for(p['id']))}}
    return job


@app.post('/api/jobs/{ident}/cancel')
def cancel_job(ident:str):
    return jobs.cancel(store,ident)


@app.get('/api/source-runs')
def source_runs():
    return store.list('source_runs')


@app.get('/api/local-sources')
def local_sources():
    captures=store.list('local_captures')
    return {'regions':REGIONS,'captures':[{k:c[k] for k in ('id','region','title','source_locator','data_period','retrieved_at','scope')} for c in captures]}


@app.post('/api/local-captures')
def import_local_capture(body:dict):
    return save_capture(store,body)


@app.post('/api/projects/{pid}/local-captures/{ident}')
def add_local_capture(pid:str,ident:str):
    project_or_404(pid)
    return use_capture(store,pid,ident)


@app.post('/api/projects/{pid}/evidence')
def add_evidence(pid:str,body:Evidence):
    project_or_404(pid)
    return store.save('evidence',{**body.model_dump(),'project_id':pid,'retrieved_at':utcnow()})


@app.post('/api/projects/{pid}/evidence/{eid}/review')
def review_evidence(pid:str,eid:str,body:Evidence):
    project_or_404(pid)
    previous=store.get('evidence',eid)
    if not previous or previous.get('project_id')!=pid: raise HTTPException(404,'证据不存在')
    original={k:v for k,v in previous.items() if k not in ('id','created_at','updated_at')}
    reviewed={**original,**body.model_dump(),'supersedes':eid,'reviewed_at':utcnow()}
    if previous.get('source_id'):
        # Verification cannot turn an external API observation into a project experiment.
        reviewed.update(source_type='market',level=min(body.level,1),source_locator=previous['source_locator'],content=previous['content'])
    return store.save('evidence',reviewed)


def parse_file(name,content):
    suffix=Path(name).suffix.lower()
    if suffix in ('.docx','.xlsx'):
        with zipfile.ZipFile(BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist())>100_000_000: raise ValueError('文件解压后过大，请拆分后上传')
    if suffix in ('.txt','.md','.csv','.json'):
        for encoding in ('utf-8-sig','gb18030'):
            try: return content.decode(encoding)[:200000]
            except UnicodeDecodeError: pass
        raise ValueError('文件编码无法识别，请另存为UTF-8文本')
    if suffix=='.docx':
        from docx import Document
        d=Document(BytesIO(content))
        return '\n'.join([p.text for p in d.paragraphs]+[' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows])[:200000]
    if suffix=='.xlsx':
        from openpyxl import load_workbook
        w=load_workbook(BytesIO(content),read_only=True,data_only=True)
        lines=[]
        for sheet in w:
            lines.append(sheet.title)
            for row in sheet.iter_rows(max_row=5000,values_only=True): lines.append(' | '.join(str(v) if v is not None else '' for v in row))
        w.close()
        return '\n'.join(lines)[:200000]
    if suffix=='.pdf':
        from pypdf import PdfReader
        return '\n'.join(p.extract_text() or '' for p in PdfReader(BytesIO(content)).pages[:100])[:200000]
    raise ValueError('支持 TXT、MD、CSV、JSON、DOCX、XLSX 和可复制文字的PDF')


@app.post('/api/projects/{pid}/upload')
def upload(pid:str,file:UploadFile):
    project_or_404(pid)
    content=file.file.read(20_000_001)
    if len(content)>20_000_000: raise ValueError('文件超过20MB，请拆分后上传')
    try: text=parse_file(file.filename or '',content)
    except ValueError: raise
    except Exception: raise ValueError('文件无法解析，请检查文件是否损坏或受密码保护') from None
    if not text.strip(): raise ValueError('未提取到文字；扫描文件请先转成可复制的文字')
    return add_evidence(pid,Evidence(title=file.filename or '上传资料',source_locator='上传文件：'+(file.filename or '资料'),content=text))


@app.post('/api/projects/{pid}/attachments',status_code=202)
def attachment(pid:str,file:UploadFile, label:str=Form('')):
    project_or_404(pid)
    from sabc.attachments import extract, IMAGE_SUFFIXES, DOCUMENT_SUFFIXES
    name=Path(file.filename or '附件').name
    if Path(name).suffix.lower() not in IMAGE_SUFFIXES|DOCUMENT_SUFFIXES:
        raise ValueError('此格式无法读取，请上传常见文档、图片或先在浏览器抽取视频画面')
    content=file.file.read(20_000_001)
    if len(content)>20_000_000: raise ValueError('单个文档最大20MB；图片和视频请先完成本地处理')
    ident=str(uuid4())
    directory=store.path.parent/'attachments';directory.mkdir(exist_ok=True)
    path=directory/ident
    path.write_bytes(content);path.chmod(0o600)
    def process():
        config=settings()
        result=extract(name,content,config,model_key(config),parse_file)
        with jobs.lock:
            check_cancelled()
            return store.save('evidence',{'project_id':pid,'title':(label or name)[:200],
                'source_locator':'对话附件：'+name,'attachment_id':ident,'source_type':'user',
                'verification_status':'unverified','level':0,'retrieved_at':utcnow(),
                'data_period':'待核对','scope':label[:500] or '本项目附件，适用范围待核对',**result})
    try:
        return jobs.submit(store,ident,pid,{'operation':'attachment','name':name},process)
    except Exception:
        path.unlink(missing_ok=True)
        raise


@app.post('/api/company/import')
def import_company(file:UploadFile):
    content=file.file.read(20_000_001)
    if len(content)>20_000_000: raise ValueError('文件超过20MB')
    try: return {'text':parse_file(file.filename or '',content),'note':'请核对并填写公司资料，导入内容不会自动成为已确认事实。'}
    except ValueError: raise
    except Exception: raise ValueError('文件无法解析') from None


@app.post('/api/projects/{pid}/assess')
def evaluate(pid:str,body:dict):
    p=project_or_404(pid)
    life=lifecycle.state(p)
    if p.get('lifecycle') and not life['confirmed']:
        raise ValueError('请先确认项目实际阶段，再生成阶段评级报告')
    if any(p.get(k)!=v for k,v in p.get('pending_patch',{}).items()):
        raise ValueError('模型整理了待核对的项目事实，请先到项目资料核对并保存，再生成评级')
    proposal=body.get('proposal') or p.get('proposal') or {}
    if proposal and body.get('confirmed') is not True: raise ValueError('请先核对并确认评分建议')
    return persist_assessment(build_assessment(p, company(), evidence_for(pid), validate_proposal(proposal)))


def build_assessment(project, c, e, proposal):
    from sabc.standard import reconcile_unknown_scores
    p=deepcopy(project)
    proposal=deepcopy(proposal)
    reconcile_unknown_scores(proposal)
    report_grounding.apply_decision_facts(p, proposal)
    p['proposal']=proposal
    p.pop('assessment_review', None)
    life=lifecycle.state(p)
    result=assess(p,c,e,proposal)
    if result['grade'] == 'NR':
        current_gaps = proposal.get('grounding_version') == report_grounding.VERSION
        explained = {dim['name'] + '尚无法判断' for dim in result['dimensions'] if dim['score'] is None} if current_gaps else set()
        missing = '、'.join(item for item in result['missing'] if item not in explained)
        details = gap_details(p, result['dimensions'] if current_gaps else None)
        result['deferral_reason'] = ('暂缓评级：以下信息还不足以支持判断。' + ('还需确认：' + missing + '。' if missing else '')
                                    + ('\n' + '\n'.join(details) if details else '')
                                    + '\n先补齐影响判断的资料，再评等级；资料缺失不代表事实不存在。')
    if p.get('lifecycle'):
        result.update(stage=life['stage'], provisional=True,
                      status='待评级' if result['grade']=='NR' else '阶段暂定评级')
    result.update(business_stage=p.get('framing', {}).get('business_stage', 'unknown'),
                  assessment_purpose=p.get('framing', {}).get('purpose', 'unknown'))
    # A generated recommendation must not override the engine's final decision.
    if result['grade'] in ('C', 'NR') and life.get('review'):
        deferred = result['grade'] == 'NR'
        negative_reasons = [d['reason'] for d in result['dimensions'] if d['score'] is not None and d['score'] < 3]
        explanation = result.get('deferral_reason') if deferred else (
            '当前不立项：' + '；'.join(negative_reasons or result['triggered_rules'] or [result['action']]))
        purpose=p.get('framing', {}).get('purpose')
        if purpose == 'research':
            next_action=('补齐影响判断的公开或经授权资料，再重新评估；现有材料不足以支持新增投入的判断。' if deferred else
                         '核查导致低分或否决的实际条件及其改善情况，再重新评估。')
            if not deferred:
                explanation='当前条件不支持该方案：' + '；'.join(negative_reasons or result['triggered_rules'] or [result['action']])
        else:
            next_action=('先补齐影响本次判断的关键依据，再重新评估；依据不足时不建议新增投入。' if deferred else
                         '先调整导致低分或否决的实际条件，再重新评估；不继续原方案新增投入。')
        life['review'] = {**life['review'], 'conclusion': 'needs_info' if deferred else 'not_recommended',
                          'summary': explanation, 'next_action': next_action}
        p['lifecycle'] = life
        if proposal.get('decision_brief'):
            proposal['decision_brief']['allocation'] = next_action
    elif life.get('review') and life['review'].get('summary', '').startswith('暂缓评级：'):
        life['review'] = {**life['review'],
                          'summary': '仍需验证的依据：' + life['review']['summary'][5:]}
        p['lifecycle'] = life
    # The latest entry belongs to this candidate; older saved report snapshots stay immutable.
    if life.get('review') and life.get('reviews') and life['reviews'][-1].get('time') == life['review'].get('time'):
        life['reviews'][-1] = deepcopy(life['review'])
        p['lifecycle'] = life
    return {'project_id':p['id'],'result':result,
            'snapshot':{'project':p,'company':deepcopy(c),'evidence':deepcopy(e),'proposal':proposal}}


def persist_assessment(candidate):
    record=store.save('assessments',candidate)
    p=deepcopy(candidate['snapshot']['project'])
    result=candidate['result']
    if result['grade']!='NR' and p.get('lifecycle', {}).get('stage')=='post':
        p['lifecycle']['next_review_on']=None
    store.save('projects',{**p,'proposal':candidate['snapshot']['proposal'],'last_grade':result['grade'],'last_assessment_id':record['id']})
    return record


@app.get('/api/assessments/{aid}/export')
def export(aid:str):
    record=store.get('assessments',aid)
    if not record: raise HTTPException(404,'报告不存在')
    return Response(json.dumps(record,ensure_ascii=False,indent=2),media_type='application/json',
                    headers={'Content-Disposition':f'attachment; filename="sabc-{aid[:8]}.json"'})


@app.put('/api/settings')
def save_settings(body:dict):
    if sso.enabled(): raise HTTPException(403,'模型由服务端统一配置')
    from urllib.parse import urlparse
    base=str(body.get('base_url','')).strip().rstrip('/')
    if base and (urlparse(base).scheme not in ('https','http') or urlparse(base).username): raise ValueError('模型地址须为不含用户名密码的HTTP(S)地址')
    if base and not urlparse(base).hostname: raise ValueError('请填写完整的模型服务地址')
    if base and os.getenv('SABC_UI_ORIGIN','').startswith('https://') and urlparse(base).scheme != 'https':
        raise ValueError('云端模型连接必须使用HTTPS地址')
    s=settings()
    has_key=not s.get('key_disabled') and bool(s.get('encrypted_key') or os.getenv('SABC_API_KEY'))
    if key_origin(base) != key_origin(s.get('base_url','')) and has_key and not (body.get('api_key') or body.get('clear_key')):
        raise ValueError('模型地址已改变，请为新服务配置对应密钥，或明确清除原密钥')
    data={'id':'model','base_url':base,'model':str(body.get('model','')).strip(),'encrypted_key':s.get('encrypted_key',''), 'key_disabled':s.get('key_disabled',False)}
    if body.get('api_key'): data.update(encrypted_key=encrypt_key(body['api_key']),key_disabled=False)
    if body.get('clear_key'): data.update(encrypted_key='',key_disabled=True)
    store.save('settings',data)
    return public_settings()
